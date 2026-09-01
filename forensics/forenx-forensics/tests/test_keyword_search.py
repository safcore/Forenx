"""Utilities and tests for Phase 4 keyword search."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.services.keyword_service import KeywordSearchService
from app.utils.exceptions import (
    EvidenceFileError,
    KeywordSearchError,
    UnsupportedFileTypeError,
)


def _write_simple_pdf(path: Path, text: str) -> None:
    """Write a minimal one-page PDF containing ``text``."""
    escaped = (
        text.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET\n"
    stream_bytes = stream.encode("latin-1", errors="replace")
    objects = [
        "1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n",
        "2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n",
        (
            "3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Contents 4 0 R /Resources<< /Font<< /F1 5 0 R >> >> >>endobj\n"
        ),
        (
            f"4 0 obj<< /Length {len(stream_bytes)} >>stream\n".encode("ascii")
            + stream_bytes
            + b"endstream\nendobj\n"
        ),
        "5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj\n",
    ]

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(output))
        if isinstance(obj, bytes):
            output.extend(obj)
        else:
            output.extend(obj.encode("ascii"))

    xref_pos = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer<< /Size {len(offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF\n"
        ).encode("ascii")
    )
    path.write_bytes(bytes(output))


@pytest.fixture
def service(keyword_service: KeywordSearchService) -> KeywordSearchService:
    """Alias keyword_service for local readability."""
    return keyword_service


@pytest.fixture
def text_dir(tmp_path: Path) -> Path:
    """Create a directory of mixed text evidence samples."""
    (tmp_path / "notes.txt").write_text(
        "Line one mentions confidential data.\n"
        "Line two has a Password reset token.\n"
        "Line three talks about bitcoin wallet theft.\n"
        "Unicode café and 机密文件 appear here.\n",
        encoding="utf-8",
    )
    (tmp_path / "access.log").write_text(
        "INFO upload dropbox confidential.pdf\n"
        "WARN usb device mounted by employee\n"
        "ERROR malware detected on endpoint\n",
        encoding="utf-8",
    )
    (tmp_path / "export.csv").write_text(
        "id,note\n1,salary adjustment for Project X\n2,no keyword here\n",
        encoding="utf-8",
    )
    (tmp_path / "case.json").write_text(
        '{"finding":"employee accessed confidential payroll","tag":"salary"}',
        encoding="utf-8",
    )
    (tmp_path / "empty.txt").write_text("", encoding="utf-8")
    (tmp_path / "binary.bin").write_bytes(b"\x00\x01\x02\x03password\x00malware")
    (tmp_path / "image.png").write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR" + b"\x00" * 32
    )
    return tmp_path


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    """Create a searchable PDF containing forensic keywords."""
    path = tmp_path / "evidence.pdf"
    _write_simple_pdf(
        path,
        "Employee uploaded confidential files to Dropbox with bitcoin wallet notes",
    )
    return path


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    """Create a searchable DOCX containing forensic keywords."""
    path = tmp_path / "evidence.docx"
    document = Document()
    document.add_paragraph("Investigation summary for Project X.")
    document.add_paragraph(
        "An employee copied confidential payroll and salary files to USB storage."
    )
    document.add_paragraph("No crypto references in this paragraph.")
    document.save(path)
    return path


class TestTextSearch:
    """Text-like evidence search tests."""

    def test_txt_search(self, service: KeywordSearchService, text_dir: Path) -> None:
        """TXT files return line-scoped matches."""
        result = service.search(text_dir / "notes.txt", "confidential")
        assert result.match_count >= 1
        assert result.matches[0].line_number == 1
        assert "confidential" in result.matches[0].matched_text.lower()
        assert "..." in result.matches[0].context or "confidential" in result.matches[0].context

    def test_log_search(self, service: KeywordSearchService, text_dir: Path) -> None:
        """LOG files are searched as text."""
        result = service.search(text_dir / "access.log", "dropbox")
        assert result.match_count == 1
        assert result.file_type == "log"

    def test_csv_search(self, service: KeywordSearchService, text_dir: Path) -> None:
        """CSV rows are searched line-by-line."""
        result = service.search(text_dir / "export.csv", "salary")
        assert result.match_count == 1
        assert result.matches[0].line_number == 2

    def test_json_search(self, service: KeywordSearchService, text_dir: Path) -> None:
        """JSON content is searchable as text."""
        result = service.search(text_dir / "case.json", "employee")
        assert result.match_count == 1

    def test_multiple_keywords(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Multiple keywords accumulate matches."""
        result = service.search_multiple(
            text_dir / "notes.txt",
            ["confidential", "bitcoin", "wallet"],
        )
        assert result.match_count >= 3
        found = {match.keyword for match in result.matches}
        assert {"confidential", "bitcoin", "wallet"} <= found

    def test_case_insensitive_default(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Matching is case-insensitive by default."""
        result = service.search(text_dir / "notes.txt", "PASSWORD")
        assert result.match_count == 1

    def test_case_sensitive_mode(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Case-sensitive mode rejects differing case."""
        result = service.search(
            text_dir / "notes.txt",
            "Password",
            case_sensitive=True,
        )
        assert result.match_count == 1
        missing = service.search(
            text_dir / "notes.txt",
            "PASSWORD",
            case_sensitive=True,
        )
        assert missing.match_count == 0
        assert missing.status.value == "no_matches"

    def test_whole_word_mode(
        self,
        service: KeywordSearchService,
        tmp_path: Path,
    ) -> None:
        """Whole-word mode avoids partial token matches."""
        path = tmp_path / "words.txt"
        path.write_text("malwaredemo and malware family\n", encoding="utf-8")
        partial = service.search(path, "malware", whole_word=False)
        whole = service.search(path, "malware", whole_word=True)
        assert partial.match_count >= 2
        assert whole.match_count == 1

    def test_partial_word_default(
        self,
        service: KeywordSearchService,
        tmp_path: Path,
    ) -> None:
        """Default mode allows partial-word matches."""
        path = tmp_path / "partial.txt"
        path.write_text("confidentiality agreement\n", encoding="utf-8")
        result = service.search(path, "confidential")
        assert result.match_count == 1

    def test_regex_search(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Regex mode supports patterned matching."""
        result = service.search(
            text_dir / "notes.txt",
            r"bit+coin",
            regex=True,
        )
        assert result.match_count == 1

    def test_invalid_regex(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Invalid regex patterns raise KeywordSearchError."""
        with pytest.raises(KeywordSearchError, match="Invalid regular expression"):
            service.search(text_dir / "notes.txt", "(", regex=True)

    def test_unicode_search(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Unicode keywords are supported in UTF-8 text files."""
        result = service.search(text_dir / "notes.txt", "café")
        assert result.match_count == 1
        result_cjk = service.search(text_dir / "notes.txt", "机密文件")
        assert result_cjk.match_count == 1

    def test_empty_file(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Empty files return zero matches without error."""
        result = service.search(text_dir / "empty.txt", "confidential")
        assert result.match_count == 0
        assert result.status.value == "no_matches"

    def test_missing_file(self, service: KeywordSearchService, tmp_path: Path) -> None:
        """Missing files raise EvidenceFileError."""
        with pytest.raises(EvidenceFileError):
            service.search(tmp_path / "missing.txt", "secret")

    def test_unsupported_file(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Unsupported formats raise UnsupportedFileTypeError."""
        with pytest.raises(UnsupportedFileTypeError):
            service.search(text_dir / "image.png", "secret")


class TestPdfAndDocxSearch:
    """PDF and DOCX searcher tests."""

    def test_pdf_search(self, service: KeywordSearchService, pdf_file: Path) -> None:
        """PDF matches include page numbers and context."""
        result = service.search(pdf_file, "confidential")
        assert result.match_count >= 1
        assert result.matches[0].page_number == 1
        assert "confidential" in result.matches[0].context.lower()

    def test_docx_search(self, service: KeywordSearchService, docx_file: Path) -> None:
        """DOCX matches include paragraph numbers."""
        result = service.search(docx_file, "confidential")
        assert result.match_count == 1
        assert result.matches[0].paragraph_number == 2
        assert "salary" in result.matches[0].context.lower()
        assert result.matches[0].matched_text.lower() == "confidential"


class TestDirectoryAndSummary:
    """Directory search and summary tests."""

    def test_directory_search_and_summary(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Directory search skips unsupported/binary files and summarizes."""
        response = service.search_directory(
            text_dir,
            ["confidential", "malware", "salary"],
        )
        assert response.summary.total_files_searched >= 5
        assert response.summary.total_matches >= 1
        assert response.summary.files_with_matches >= 1
        assert response.summary.files_skipped >= 1
        assert "confidential" in response.summary.keyword_frequency
        assert response.summary.execution_time_ms >= 0

    def test_binary_file_skipped_in_directory(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """Binary files are skipped during directory search."""
        # Rename binary to a text-like extension to exercise binary detection.
        binary_as_txt = text_dir / "payload.txt"
        binary_as_txt.write_bytes(b"\x00\x00\x00\x00\xff\xfe malware confidential")
        response = service.search_directory(text_dir, ["malware"])
        skipped = [item for item in response.results if item.file_name == "payload.txt"]
        assert skipped
        assert skipped[0].status.value == "skipped"

    def test_summarize_results_helper(
        self,
        service: KeywordSearchService,
        text_dir: Path,
    ) -> None:
        """summarize_results aggregates per-file outcomes."""
        first = service.search(text_dir / "notes.txt", "bitcoin")
        second = service.search(text_dir / "empty.txt", "bitcoin")
        summary = service.summarize_results([first, second])
        assert summary.total_files_searched == 2
        assert summary.files_with_matches == 1
        assert summary.files_without_matches == 1
        assert summary.total_matches == first.match_count
